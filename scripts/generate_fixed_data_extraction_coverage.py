from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE_COVERAGE_JSON = ROOT / "data" / "catalogue" / "fixed_data_source_coverage.json"
MANUAL_EXTRACTION_INDEX_JSON = ROOT / "data" / "catalogue" / "manual_extraction_index.json"
MANUAL_EXTRACTION_ROOT = ROOT / "data" / "catalogue" / "manual_extractions"
OUTPUT_JSON = ROOT / "data" / "catalogue" / "fixed_data_extraction_coverage.json"
OUTPUT_XLSX = ROOT / "Fixed_Data_Extraction_Coverage.xlsx"


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def extract_system_id(path_text: str) -> str:
    match = re.search(r"(SYS_[A-Z0-9_]+)_extraction\.json$", path_text)
    return match.group(1) if match else ""


def autosize_columns(worksheet) -> None:
    for column in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            value = normalize_text(cell.value)
            max_length = max(max_length, len(value))
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 48)


def style_sheet(worksheet) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="DCE6F1")
    header_font = Font(bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    autosize_columns(worksheet)


def load_extraction_details(index_rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    details: dict[str, dict[str, Any]] = {}
    for row in index_rows:
        extraction_file = normalize_text(row.get("Extraction File", ""))
        system_id = extract_system_id(extraction_file)
        if not system_id:
            continue
        payload_path = ROOT / extraction_file
        payload = load_json(payload_path) if payload_path.exists() else {}
        systems = payload.get("systems", []) or []
        trays = payload.get("trays", []) or []
        components = payload.get("components", []) or []
        products = payload.get("products", []) or []
        mappings = payload.get("system_mappings", []) or []
        unresolved = payload.get("unresolved_items", []) or []
        source_meta = payload.get("source_meta", {}) or {}
        has_description = any(normalize_text(item.get("system_notes", "")) for item in systems)
        details[system_id] = {
            "Supplier Name": normalize_text(row.get("Supplier Name", "")),
            "System Name": normalize_text(row.get("System Name", "")),
            "Extraction File": extraction_file,
            "Source Type": normalize_text(row.get("Source Type", "")),
            "Source Document": normalize_text(row.get("Source Document", "")),
            "Source Version": normalize_text(row.get("Source Version", "")),
            "Source Date": normalize_text(row.get("Source Date", "")),
            "Systems Count": len(systems),
            "Trays Count": len(trays),
            "Components Count": len(components),
            "Products Count": len(products),
            "Mappings Count": len(mappings),
            "Unresolved Count": len(unresolved),
            "Has Description": "yes" if has_description else "no",
            "Source Meta System Name": normalize_text(source_meta.get("system_name", "")),
        }
    return details


def summarize_group(rows: list[dict[str, Any]], group_key: str) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = defaultdict(
        lambda: {
            "Systems Total": 0,
            "Systems Extracted": 0,
            "Systems Missing Extraction": 0,
            "Systems With Descriptions": 0,
            "Components Total": 0,
            "Products Total": 0,
            "Mappings Total": 0,
            "Unresolved Total": 0,
        }
    )

    for row in rows:
        label = normalize_text(row.get(group_key, "")) or "Unspecified"
        bucket = grouped[label]
        bucket["Systems Total"] += 1
        extracted = row["Extraction Status"] == "extracted"
        if extracted:
            bucket["Systems Extracted"] += 1
            bucket["Systems With Descriptions"] += 1 if row["Has Description"] == "yes" else 0
            bucket["Components Total"] += int(row["Components Count"])
            bucket["Products Total"] += int(row["Products Count"])
            bucket["Mappings Total"] += int(row["Mappings Count"])
            bucket["Unresolved Total"] += int(row["Unresolved Count"])
        else:
            bucket["Systems Missing Extraction"] += 1

    summary_rows: list[dict[str, Any]] = []
    label_field = group_key
    for label, bucket in sorted(
        grouped.items(),
        key=lambda item: (-item[1]["Systems Missing Extraction"], -item[1]["Systems Total"], item[0]),
    ):
        total = bucket["Systems Total"]
        extracted = bucket["Systems Extracted"]
        summary_rows.append(
            {
                label_field: label,
                **bucket,
                "Coverage %": round((extracted / total) * 100, 1) if total else 0.0,
            }
        )
    return summary_rows


def build_rows() -> tuple[list[dict[str, Any]], dict[str, Any], list[dict[str, Any]], list[dict[str, Any]]]:
    source_coverage = load_json(SOURCE_COVERAGE_JSON)
    extraction_index = load_json(MANUAL_EXTRACTION_INDEX_JSON)

    extracted_by_system = load_extraction_details(extraction_index.get("rows", []) or [])
    systems_rows = source_coverage.get("system_coverage", []) or []

    rows: list[dict[str, Any]] = []
    extracted_count = 0
    missing_count = 0
    systems_with_components = 0
    systems_with_products = 0
    systems_with_mappings = 0
    systems_with_unresolved = 0
    systems_with_descriptions = 0
    total_components = 0
    total_products = 0
    total_mappings = 0
    total_unresolved = 0

    for system in systems_rows:
        system_id = normalize_text(system.get("System ID", ""))
        extraction = extracted_by_system.get(system_id)
        extracted = extraction is not None
        if extracted:
            extracted_count += 1
        else:
            missing_count += 1

        row = {
            "Supplier Name": normalize_text(system.get("Supplier Name", "")),
            "System ID": system_id,
            "System Name": normalize_text(system.get("System Name", "")),
            "System Category": normalize_text(system.get("System Category", "")),
            "Source Count": normalize_text(system.get("Source Count", "")),
            "Extraction Status": "extracted" if extracted else "missing",
            "Extraction File": extraction.get("Extraction File", "") if extraction else "",
            "Extraction Source Type": extraction.get("Source Type", "") if extraction else "",
            "Systems Count": extraction.get("Systems Count", 0) if extraction else 0,
            "Trays Count": extraction.get("Trays Count", 0) if extraction else 0,
            "Components Count": extraction.get("Components Count", 0) if extraction else 0,
            "Products Count": extraction.get("Products Count", 0) if extraction else 0,
            "Mappings Count": extraction.get("Mappings Count", 0) if extraction else 0,
            "Unresolved Count": extraction.get("Unresolved Count", 0) if extraction else 0,
            "Has Description": extraction.get("Has Description", "no") if extraction else "no",
            "Source Folder": normalize_text(system.get("Source Folder", "")),
        }
        rows.append(row)

        if not extracted:
            continue
        systems_with_components += 1 if row["Components Count"] else 0
        systems_with_products += 1 if row["Products Count"] else 0
        systems_with_mappings += 1 if row["Mappings Count"] else 0
        systems_with_unresolved += 1 if row["Unresolved Count"] else 0
        systems_with_descriptions += 1 if row["Has Description"] == "yes" else 0
        total_components += int(row["Components Count"])
        total_products += int(row["Products Count"])
        total_mappings += int(row["Mappings Count"])
        total_unresolved += int(row["Unresolved Count"])

    summary = {
        "systems_total": len(rows),
        "systems_extracted": extracted_count,
        "systems_missing_extraction": missing_count,
        "systems_with_components": systems_with_components,
        "systems_with_products": systems_with_products,
        "systems_with_mappings": systems_with_mappings,
        "systems_with_unresolved": systems_with_unresolved,
        "systems_with_descriptions": systems_with_descriptions,
        "components_total": total_components,
        "products_total": total_products,
        "mappings_total": total_mappings,
        "unresolved_total": total_unresolved,
    }
    supplier_coverage = summarize_group(rows, "Supplier Name")
    category_coverage = summarize_group(rows, "System Category")
    return rows, summary, supplier_coverage, category_coverage


def write_json(
    rows: list[dict[str, Any]],
    summary: dict[str, Any],
    supplier_coverage: list[dict[str, Any]],
    category_coverage: list[dict[str, Any]],
) -> None:
    payload = {
        "summary": summary,
        "supplier_coverage": supplier_coverage,
        "category_coverage": category_coverage,
        "rows": rows,
    }
    OUTPUT_JSON.write_text(json.dumps(payload, indent=2, ensure_ascii=True), encoding="utf-8")


def append_sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]], fallback_headers: list[str]) -> None:
    worksheet = workbook.create_sheet(title)
    headers = list(rows[0].keys()) if rows else fallback_headers
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header, "") for header in headers])
    style_sheet(worksheet)


def write_xlsx(
    rows: list[dict[str, Any]],
    summary: dict[str, Any],
    supplier_coverage: list[dict[str, Any]],
    category_coverage: list[dict[str, Any]],
) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "SUMMARY"
    summary_sheet.append(["Metric", "Value"])
    for key, value in summary.items():
        summary_sheet.append([key, value])
    style_sheet(summary_sheet)

    append_sheet(
        workbook,
        "SUPPLIER_COVERAGE",
        supplier_coverage,
        [
            "Supplier Name",
            "Systems Total",
            "Systems Extracted",
            "Systems Missing Extraction",
            "Systems With Descriptions",
            "Components Total",
            "Products Total",
            "Mappings Total",
            "Unresolved Total",
            "Coverage %",
        ],
    )

    append_sheet(
        workbook,
        "CATEGORY_COVERAGE",
        category_coverage,
        [
            "System Category",
            "Systems Total",
            "Systems Extracted",
            "Systems Missing Extraction",
            "Systems With Descriptions",
            "Components Total",
            "Products Total",
            "Mappings Total",
            "Unresolved Total",
            "Coverage %",
        ],
    )

    headers = list(rows[0].keys()) if rows else [
        "Supplier Name",
        "System ID",
        "System Name",
        "System Category",
        "Source Count",
        "Extraction Status",
        "Extraction File",
        "Extraction Source Type",
        "Systems Count",
        "Trays Count",
        "Components Count",
        "Products Count",
        "Mappings Count",
        "Unresolved Count",
        "Has Description",
        "Source Folder",
    ]
    append_sheet(workbook, "SYSTEM_COVERAGE", rows, headers)

    missing_rows = [row for row in rows if row["Extraction Status"] == "missing"]
    append_sheet(workbook, "MISSING_EXTRACTIONS", missing_rows, headers)

    workbook.save(OUTPUT_XLSX)


def main() -> None:
    rows, summary, supplier_coverage, category_coverage = build_rows()
    write_json(rows, summary, supplier_coverage, category_coverage)
    write_xlsx(rows, summary, supplier_coverage, category_coverage)
    print(OUTPUT_JSON)
    print(OUTPUT_XLSX)


if __name__ == "__main__":
    main()
