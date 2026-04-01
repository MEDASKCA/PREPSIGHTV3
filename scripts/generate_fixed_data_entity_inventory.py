from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DATASET_JSON = ROOT / "data" / "catalogue" / "fixed_data_dataset_batch668.json"
DEFAULT_OUTPUT_JSON = ROOT / "data" / "catalogue" / "fixed_data_entity_inventory.json"
DEFAULT_OUTPUT_XLSX = ROOT / "Fixed_Data_Entity_Inventory.xlsx"


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", normalize_text(value).lower()).strip()


def autosize_columns(worksheet) -> None:
    for column in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            max_length = max(max_length, len(normalize_text(cell.value)))
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 48)


def style_sheet(worksheet) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="DCE6F1")
    header_font = Font(bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    autosize_columns(worksheet)


def append_sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]], fallback_headers: list[str]) -> None:
    worksheet = workbook.create_sheet(title)
    headers = list(rows[0].keys()) if rows else fallback_headers
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header, "") for header in headers])
    style_sheet(worksheet)


def build_inventory(dataset: dict[str, Any]) -> dict[str, Any]:
    suppliers = dataset["SUPPLIERS"]["rows"]
    systems = dataset["SYSTEMS"]["rows"]
    trays = dataset["TRAYS"]["rows"]
    components = dataset["COMPONENTS"]["rows"]
    products = dataset["PRODUCTS"]["rows"]
    mappings = dataset["SYSTEM_MAPPINGS"]["rows"]

    supplier_by_id = {normalize_text(row.get("Supplier ID", "")): row for row in suppliers}
    system_by_id = {normalize_text(row.get("System ID", "")): row for row in systems}
    component_by_id = {normalize_text(row.get("Component ID", "")): row for row in components}

    trays_by_system = Counter(normalize_text(row.get("System ID", "")) for row in trays)
    components_by_system = Counter(normalize_text(row.get("System ID", "")) for row in components)
    mappings_by_system = Counter(normalize_text(row.get("System ID", "")) for row in mappings)
    products_by_component = Counter(normalize_text(row.get("Component ID", "")) for row in products)

    product_count_by_system: Counter[str] = Counter()
    for component_id, count in products_by_component.items():
        system_id = normalize_text(component_by_id.get(component_id, {}).get("System ID", ""))
        if system_id:
            product_count_by_system[system_id] += count

    system_rows: list[dict[str, Any]] = []
    for row in systems:
        system_id = normalize_text(row.get("System ID", ""))
        supplier_id = normalize_text(row.get("Supplier ID", ""))
        supplier_name = normalize_text(supplier_by_id.get(supplier_id, {}).get("Supplier Name", ""))
        description = normalize_text(row.get("System Notes", ""))
        system_rows.append(
            {
                "System ID": system_id,
                "System Name": normalize_text(row.get("System Name", "")),
                "Supplier Name": supplier_name,
                "System Category": normalize_text(row.get("System Category", "")),
                "Has Description": "yes" if description else "no",
                "Description Length": len(description),
                "Tray Count": trays_by_system.get(system_id, 0),
                "Component Count": components_by_system.get(system_id, 0),
                "Product Count": product_count_by_system.get(system_id, 0),
                "Mapping Count": mappings_by_system.get(system_id, 0),
                "System Notes": description,
            }
        )

    duplicate_system_names_map: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in system_rows:
        duplicate_system_names_map[normalize_key(row["System Name"])].append(row)
    duplicate_system_names = []
    for key, items in sorted(duplicate_system_names_map.items()):
        if key and len(items) > 1:
            duplicate_system_names.append(
                {
                    "Normalized Name": key,
                    "Occurrences": len(items),
                    "Suppliers": " | ".join(sorted({row["Supplier Name"] for row in items if row["Supplier Name"]})),
                    "System IDs": " | ".join(row["System ID"] for row in items),
                    "System Names": " | ".join(row["System Name"] for row in items),
                }
            )

    component_rows: list[dict[str, Any]] = []
    for row in components:
        system_id = normalize_text(row.get("System ID", ""))
        system_name = normalize_text(system_by_id.get(system_id, {}).get("System Name", ""))
        supplier_id = normalize_text(system_by_id.get(system_id, {}).get("Supplier ID", ""))
        supplier_name = normalize_text(supplier_by_id.get(supplier_id, {}).get("Supplier Name", ""))
        component_rows.append(
            {
                "Component ID": normalize_text(row.get("Component ID", "")),
                "Component Name": normalize_text(row.get("Component Name", "")),
                "Component Role": normalize_text(row.get("Component Role", "")),
                "Implant Category": normalize_text(row.get("Implant Category", "")),
                "System ID": system_id,
                "System Name": system_name,
                "Supplier Name": supplier_name,
                "Product Count": products_by_component.get(normalize_text(row.get("Component ID", "")), 0),
                "Component Notes": normalize_text(row.get("Component Notes", "")),
            }
        )

    duplicate_component_names_map: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in component_rows:
        duplicate_component_names_map[normalize_key(row["Component Name"])].append(row)
    duplicate_component_names = []
    for key, items in sorted(duplicate_component_names_map.items()):
        if key and len(items) > 1:
            duplicate_component_names.append(
                {
                    "Normalized Name": key,
                    "Occurrences": len(items),
                    "Distinct Systems": len({row["System ID"] for row in items if row["System ID"]}),
                    "Suppliers": " | ".join(sorted({row["Supplier Name"] for row in items if row["Supplier Name"]})),
                    "Component IDs": " | ".join(row["Component ID"] for row in items),
                    "Component Names": " | ".join(row["Component Name"] for row in items),
                }
            )

    systems_with_descriptions = sum(1 for row in system_rows if row["Has Description"] == "yes")
    summary = {
        "systems_total": len(system_rows),
        "systems_with_descriptions": systems_with_descriptions,
        "systems_without_descriptions": len(system_rows) - systems_with_descriptions,
        "trays_total": len(trays),
        "components_total": len(component_rows),
        "products_total": len(products),
        "system_mappings_total": len(mappings),
        "duplicate_system_names": len(duplicate_system_names),
        "duplicate_component_names": len(duplicate_component_names),
        "alias_fields_present_in_fixed_data_workbook": "no",
    }

    return {
        "summary": summary,
        "system_entities": system_rows,
        "component_entities": component_rows,
        "duplicate_system_names": duplicate_system_names,
        "duplicate_component_names": duplicate_component_names,
    }


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=True), encoding="utf-8")


def write_xlsx(path: Path, payload: dict[str, Any]) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "SUMMARY"
    summary_sheet.append(["Metric", "Value"])
    for key, value in payload["summary"].items():
        summary_sheet.append([key, value])
    style_sheet(summary_sheet)

    append_sheet(
        workbook,
        "SYSTEM_ENTITIES",
        payload["system_entities"],
        [
            "System ID",
            "System Name",
            "Supplier Name",
            "System Category",
            "Has Description",
            "Description Length",
            "Tray Count",
            "Component Count",
            "Product Count",
            "Mapping Count",
            "System Notes",
        ],
    )
    append_sheet(
        workbook,
        "COMPONENT_ENTITIES",
        payload["component_entities"],
        [
            "Component ID",
            "Component Name",
            "Component Role",
            "Implant Category",
            "System ID",
            "System Name",
            "Supplier Name",
            "Product Count",
            "Component Notes",
        ],
    )
    append_sheet(
        workbook,
        "DUP_SYSTEM_NAMES",
        payload["duplicate_system_names"],
        ["Normalized Name", "Occurrences", "Suppliers", "System IDs", "System Names"],
    )
    append_sheet(
        workbook,
        "DUP_COMPONENT_NAMES",
        payload["duplicate_component_names"],
        ["Normalized Name", "Occurrences", "Distinct Systems", "Suppliers", "Component IDs", "Component Names"],
    )

    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate extracted-entity inventory and duplicate-name review artifacts.")
    parser.add_argument("--dataset", type=Path, default=DEFAULT_DATASET_JSON, help="Path to exported fixed-data dataset JSON.")
    parser.add_argument("--output-json", type=Path, default=DEFAULT_OUTPUT_JSON, help="Path to output JSON file.")
    parser.add_argument("--output-xlsx", type=Path, default=DEFAULT_OUTPUT_XLSX, help="Path to output XLSX file.")
    args = parser.parse_args()

    payload = build_inventory(load_json(args.dataset))
    write_json(args.output_json, payload)
    write_xlsx(args.output_xlsx, payload)
    print(args.output_json)
    print(args.output_xlsx)


if __name__ == "__main__":
    main()
